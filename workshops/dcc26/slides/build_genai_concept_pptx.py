"""Build an editable PowerPoint slide of the GenAI concept schematic.

Every element is a native, individually-adjustable object:
  - noise image      -> inserted picture
  - conditions chip  -> rounded rectangle + text
  - "+"              -> text box
  - 3 arrows         -> block-arrow autoshapes
  - GenAI box        -> rounded rectangle + text
  - output beam      -> inserted picture
  - all labels/title -> text boxes

Run with a python that has python-pptx (the base anaconda env):

    /opt/anaconda3/bin/python workshops/dcc26/slides/build_genai_concept_pptx.py

Requires the component PNGs from export_genai_components.py to exist first.
"""

from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.abspath(os.path.join(HERE, os.pardir, "assets"))
COMP = os.path.join(ASSETS, "genai_components")

# ---- palette (matches the deck) ----
BG = RGBColor(0xFB, 0xFA, 0xF7)
NEAR_BLACK = RGBColor(0x17, 0x21, 0x2B)
BOX_BLACK = RGBColor(0x11, 0x14, 0x18)
BLUE = RGBColor(0x22, 0x5E, 0x9B)
BLUE_TINT = RGBColor(0xEC, 0xF3, 0xFA)
GREY = RGBColor(0x5A, 0x66, 0x75)
GREY_SOFT = RGBColor(0x7A, 0x84, 0x93)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xCF, 0xD6, 0xDD)

FONT_DISPLAY = "Aptos Display"
FONT_BODY = "Aptos"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

TGT = {"vf": 0.30, "rmin": 2.0, "force": 0.50}


def style_line(shp, color=None, width=None):
    if color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = color
        if width is not None:
            shp.line.width = Pt(width)


def add_shape(slide, shape, x, y, w, h, *, fill, line=None, line_w=1.5):
    shp = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    style_line(shp, line, line_w)
    shp.shadow.inherit = False
    return shp


def set_text(shp, lines, *, anchor=MSO_ANCHOR.MIDDLE):
    """lines: list of (text, size, bold, italic, color) tuples."""
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Emu(0); tf.margin_right = Emu(0)
    tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    for i, (text, size, bold, italic, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.name = FONT_DISPLAY if bold and size >= 20 else FONT_BODY
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color


def add_textbox(slide, x, y, w, h, lines, *, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    set_text(tb, lines, anchor=anchor)
    return tb


def main() -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # background
    add_shape(slide, MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), SLIDE_W, SLIDE_H, fill=BG)

    # title
    add_textbox(slide, Inches(0.65), Inches(0.55), Inches(12.0), Inches(0.8),
                [("Generative AI for inverse design", 30, True, False, NEAR_BLACK)],
                anchor=MSO_ANCHOR.MIDDLE)

    # ---- noise image (2:1) ----
    nx, ny, nw, nh = Inches(0.85), Inches(2.05), Inches(3.1), Inches(1.55)
    pic = slide.shapes.add_picture(os.path.join(COMP, "noise.png"), nx, ny, nw, nh)
    style_line(pic, NEAR_BLACK, 1.5)
    add_textbox(slide, nx, Inches(3.66), nw, Inches(0.35),
                [("random noise  z   (50x100)", 13, False, False, NEAR_BLACK)],
                anchor=MSO_ANCHOR.TOP)

    # "+"
    add_textbox(slide, nx, Inches(4.02), nw, Inches(0.4),
                [("+", 22, True, False, GREY_SOFT)], anchor=MSO_ANCHOR.MIDDLE)

    # ---- conditions chip ----
    cx, cy, cw, ch = Inches(0.85), Inches(4.5), Inches(3.1), Inches(0.95)
    chip = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, cw, ch,
                     fill=BLUE_TINT, line=BLUE, line_w=1.75)
    set_text(chip, [
        ("conditions", 14, True, False, BLUE),
        (f"vf={TGT['vf']:.2f}    rmin={TGT['rmin']:.1f}    force@{TGT['force']:.2f}",
         12.5, False, False, NEAR_BLACK),
    ])

    # ---- arrows into the box ----
    add_shape(slide, MSO_SHAPE.RIGHT_ARROW, Inches(4.05), Inches(2.55),
              Inches(1.5), Inches(0.55), fill=NEAR_BLACK)
    add_shape(slide, MSO_SHAPE.RIGHT_ARROW, Inches(4.05), Inches(4.7),
              Inches(1.5), Inches(0.55), fill=NEAR_BLACK)

    # ---- GenAI black box ----
    bx, by, bw, bh = Inches(5.7), Inches(2.95), Inches(2.3), Inches(1.95)
    box = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, bw, bh, fill=BOX_BLACK)
    set_text(box, [
        ("GenAI", 26, True, False, WHITE),
        ("model", 15, False, False, LIGHT),
        ("(black box)", 11, False, True, GREY_SOFT),
    ])

    # ---- arrow to output ----
    add_shape(slide, MSO_SHAPE.RIGHT_ARROW, Inches(8.15), Inches(3.65),
              Inches(1.45), Inches(0.55), fill=NEAR_BLACK)

    # ---- output beam (2:1) ----
    ox, oy, ow, oh = Inches(9.75), Inches(2.9), Inches(3.0), Inches(1.5)
    pic2 = slide.shapes.add_picture(os.path.join(COMP, "beam.png"), ox, oy, ow, oh)
    style_line(pic2, NEAR_BLACK, 1.5)
    add_textbox(slide, ox, Inches(4.46), ow, Inches(0.7),
                [("new beam design", 13.5, True, False, NEAR_BLACK),
                 ("(out-of-sample conditions)", 11.5, False, True, GREY)],
                anchor=MSO_ANCHOR.TOP)

    out = os.path.join(HERE, "beams2d_genai_concept.pptx")
    prs.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
