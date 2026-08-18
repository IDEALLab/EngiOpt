"""Static 'Same number, different problem' slide with real dataset snapshots.

Standalone editable .pptx: two paper cards (A / B), each showing a real beams2d
snapshot grid that embodies that paper's setup choices, both reporting MSE=0.04.
Card numbers are aligned to the actual dataset (volfrac 0.15-0.40, rmin 1.5-4.0).

Build (base anaconda env has python-pptx); needs the snapshot PNGs first:

    conda run -n EngiBench312 python workshops/dcc26/slides/same_number_snapshots.py
    /opt/anaconda3/bin/python workshops/dcc26/slides/build_same_number_pptx.py
"""

from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.abspath(os.path.join(HERE, os.pardir, "assets"))

BG = RGBColor(0xFB, 0xFA, 0xF7)
NEAR_BLACK = RGBColor(0x17, 0x21, 0x2B)
BLUE = RGBColor(0x22, 0x5E, 0x9B)
ORANGE = RGBColor(0xC4, 0x7B, 0x20)
RED = RGBColor(0xC8, 0x48, 0x37)
GREY = RGBColor(0x5A, 0x66, 0x75)
GREY_SOFT = RGBColor(0x7A, 0x84, 0x93)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT_DISPLAY = "Aptos Display"
FONT_BODY = "Aptos"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# profiles aligned to the real dataset (must match same_number_snapshots.py)
CARDS = [
    ("paperA", BLUE, "PAPER A  —  diffusion for topology", [
        ("Material budget", "fixed  0.35"),
        ("Filter radius", "1.5 px  (sharp)"),
        ("Test split", "near training mean"),
    ]),
    ("paperB", ORANGE, "PAPER B  —  cVAE for inverse design", [
        ("Material budget", "sampled  0.15 - 0.40"),
        ("Filter radius", "3.5 px  (smeared)"),
        ("Test split", "OOD corners (force 0 / 1)"),
    ]),
]


def rect(slide, x, y, w, h, fill, *, line=None, shape=MSO_SHAPE.RECTANGLE):
    shp = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
    shp.shadow.inherit = False
    return shp


def text(slide, x, y, w, h, lines, *, font=FONT_BODY, size=14, bold=False,
         italic=False, color=NEAR_BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0); tf.margin_right = Emu(0)
    tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    items = lines if isinstance(lines, list) else [lines]
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        # item may be str or (text, kwargs-dict)
        if isinstance(item, tuple):
            s, kw = item
        else:
            s, kw = item, {}
        r = p.add_run(); r.text = s
        r.font.name = kw.get("font", font)
        r.font.size = Pt(kw.get("size", size))
        r.font.bold = kw.get("bold", bold)
        r.font.italic = kw.get("italic", italic)
        r.font.color.rgb = kw.get("color", color)
    return tb


def main() -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W; prs.slide_height = SLIDE_H
    s = prs.slides.add_slide(prs.slide_layouts[6])

    # chrome
    rect(s, Emu(0), Emu(0), SLIDE_W, SLIDE_H, BG)
    rect(s, Emu(0), Emu(0), Inches(0.12), SLIDE_H, BLUE)
    text(s, Inches(0.44), Inches(7.02), Inches(5.0), Inches(0.2),
         "Why we cannot tell", size=8.25, color=GREY)
    text(s, Inches(8.3), Inches(7.02), Inches(4.3), Inches(0.2),
         "DCC 2026 workshop  |  EngiBench + EngiOpt", size=8.25, color=GREY, align=PP_ALIGN.RIGHT)

    # eyebrow + title + subtitle
    text(s, Inches(0.65), Inches(0.48), Inches(10.0), Inches(0.25),
         "A CONCRETE CASE", size=10, bold=True, color=BLUE)
    text(s, Inches(0.65), Inches(0.78), Inches(12.0), Inches(0.7),
         "Same task name. Same headline number.", font=FONT_DISPLAY, size=26, bold=True)
    text(s, Inches(0.65), Inches(1.52), Inches(12.0), Inches(0.4),
         'Two "Beams2D" papers both report MSE = 0.04 — but trained and tested on '
         "very different slices of the same task. The snapshots are real training data.",
         size=12.5, italic=True, color=GREY)

    # cards
    y = Inches(2.10); doc_w = Inches(5.1); doc_h = Inches(4.45)
    x_a = Inches(0.85); x_b = SLIDE_W - Inches(0.85) - doc_w

    for (img, accent, label, rows), x in zip(CARDS, (x_a, x_b)):
        rect(s, x, y, doc_w, doc_h, WHITE)
        rect(s, x, y, doc_w, Inches(0.50), accent)
        text(s, x, y + Inches(0.12), doc_w, Inches(0.30),
             label, size=11.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        text(s, x, y + Inches(0.62), doc_w, Inches(0.32),
             "Reported   MSE = 0.04", size=14, bold=True, color=accent, align=PP_ALIGN.CENTER)
        # snapshot grid (3:1)
        pic_w = doc_w - Inches(0.5)
        pic = s.shapes.add_picture(os.path.join(ASSETS, f"same_number_{img}.png"),
                                   x + Inches(0.25), y + Inches(1.08), width=pic_w)
        pic.line.color.rgb = accent; pic.line.width = Pt(1.0)
        # profile rows
        ry = y + Inches(2.95)
        for k, v in rows:
            text(s, x + Inches(0.30), ry, Inches(1.9), Inches(0.3),
                 k, size=11.5, bold=True, color=NEAR_BLACK)
            text(s, x + Inches(2.25), ry, doc_w - Inches(2.55), Inches(0.3),
                 v, size=11.5, bold=True, color=accent)
            ry += Inches(0.45)

    # big != between
    text(s, Inches(5.95), Inches(3.7), Inches(1.45), Inches(1.4),
         "≠", font=FONT_DISPLAY, size=90, bold=True, color=RED,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # kicker band
    rect(s, Inches(0.65), Inches(6.42), Inches(12.05), Inches(0.48), BLUE)
    text(s, Inches(0.65), Inches(6.52), Inches(12.05), Inches(0.3),
         "Same number. Different problem.", size=13.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    out = os.path.join(HERE, "beams2d_same_number.pptx")
    prs.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
